import io
from decimal import Decimal
from pathlib import Path

import pandas as pd
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db.models import Count, Q, Sum
from django.http import FileResponse, Http404, HttpResponseRedirect
from django.shortcuts import render
from django.urls import reverse_lazy
from django.utils import timezone
from django.views.generic import CreateView, DeleteView, ListView, UpdateView
from docx import Document
from docxtpl import DocxTemplate
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

from gymApp.models import (
    BugReport,
    ClientProfile,
    Exercise,
    Gym,
    NewFeatureRequest,
    Subscription,
    SubscriptionPayment,
    WorkoutItem,
    WorkoutPlanRun,
)
from gymReports.forms import PrintTemplateForm
from gymReports.models import PrintTemplate

DEFAULT_TEMPLATE_SPECS = {
    "gyms_directory": {
        "title": "Gyms Directory",
        "fields": [
            "gym_name",
            "street",
            "city",
            "state",
            "postal_code",
            "country",
            "contact_email",
            "contact_phone",
            "equipments",
            "created_at",
        ],
    },
    "clients_overview": {
        "title": "Clients Overview",
        "fields": [
            "username",
            "email",
            "is_active",
            "verified",
            "age",
            "weight",
            "height",
            "subscription_plan",
            "active_workout_plan",
            "stripe_customer_id",
        ],
    },
    "subscription_payments": {
        "title": "Subscription Payments",
        "fields": [
            "invoice_id",
            "username",
            "subscription_id",
            "customer_id",
            "amount_major",
            "currency",
            "paid_at",
            "created_at",
        ],
    },
}


def _build_default_docx(report_key):
    spec = DEFAULT_TEMPLATE_SPECS.get(report_key)
    if not spec:
        raise Http404("Default template not available")

    doc = Document()
    doc.add_heading(spec["title"], level=1)
    doc.add_paragraph("Generated at: {{ generated_at }}")
    doc.add_paragraph("Total records: {{ total_records }}")
    doc.add_paragraph("")
    doc.add_paragraph("{% for item in records %}")
    doc.add_paragraph("— Item {{ loop.index }} —")
    for field in spec["fields"]:
        doc.add_paragraph(f"{field}: {{{{ item.{field} }}}}")
    doc.add_paragraph("")
    doc.add_paragraph("{% endfor %}")

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.read(), spec["title"]


def ensure_default_templates(report_key=None, force=False):
    keys = [report_key] if report_key else list(DEFAULT_TEMPLATE_SPECS.keys())
    for key in keys:
        active_exists = PrintTemplate.objects.filter(report_key=key, is_active=True).exists()
        PrintTemplate.objects.filter(report_key=key).exists()

        if active_exists and not force:
            continue

        content, title = _build_default_docx(key)
        filename = f"print_templates/default_{key}.docx"
        if default_storage.exists(filename):
            default_storage.delete(filename)
        stored_path = default_storage.save(filename, ContentFile(content))

        template, _ = PrintTemplate.objects.update_or_create(
            report_key=key,
            name=f"Default - {title}",
            defaults={
                "template_file": stored_path,
                "is_active": True,
                "created_by": None,
            },
        )

        PrintTemplate.objects.filter(report_key=key).exclude(pk=template.pk).update(is_active=False)


def _to_naive(dt):
    return dt.replace(tzinfo=None) if dt else None


def _df_from_records(records, columns):
    if not records:
        return pd.DataFrame(columns=columns)
    return pd.DataFrame.from_records(records, columns=columns)


def _export_xlsx(df, sheet_name, filename):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="xlsxwriter", datetime_format="yyyy-mm-dd hh:mm") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        workbook = writer.book
        worksheet = writer.sheets[sheet_name]
        header_fmt = workbook.add_format(
            {"bold": True, "bg_color": "#1F4E79", "font_color": "white", "border": 1}
        )

        for col_num, col_name in enumerate(df.columns):
            worksheet.write(0, col_num, col_name, header_fmt)

        for i, col in enumerate(df.columns):
            series = df[col].astype(str) if not df.empty else pd.Series([col])
            max_len = max([len(str(col))] + [len(val) for val in series.tolist()])
            worksheet.set_column(i, i, min(max(max_len + 2, 12), 60))

        worksheet.freeze_panes(1, 0)
        if len(df.columns) > 0:
            worksheet.autofilter(0, 0, max(len(df), 1), len(df.columns) - 1)

    output.seek(0)
    return FileResponse(
        output,
        as_attachment=True,
        filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _export_pdf(df, title, filename):
    def register_fonts():
        font_dir = Path("/usr/share/fonts/truetype/dejavu")
        regular = font_dir / "DejaVuSans.ttf"
        bold = font_dir / "DejaVuSans-Bold.ttf"

        if regular.exists():
            pdfmetrics.registerFont(TTFont("DejaVuSans", str(regular)))
        if bold.exists():
            pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", str(bold)))

        return (
            "DejaVuSans-Bold" if bold.exists() else "Helvetica-Bold",
            "DejaVuSans" if regular.exists() else "Helvetica",
        )

    header_font, body_font = register_fonts()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24
    )
    styles = getSampleStyleSheet()
    styles["Title"].fontName = header_font
    styles["Title"].textColor = colors.HexColor("#1F4E79")
    styles["Normal"].fontName = body_font
    styles["Normal"].leading = 12
    styles["Heading4"].fontName = header_font
    styles["Heading4"].textColor = colors.HexColor("#1F4E79")

    ParagraphStyle(
        "Label",
        parent=styles["Normal"],
        textColor=colors.HexColor("#1F4E79"),
        spaceAfter=2,
    )

    elements = [Paragraph(title, styles["Title"]), Spacer(1, 8)]

    if df.empty:
        elements.append(Paragraph("No data available.", styles["Normal"]))
    else:
        for idx, row in df.iterrows():
            elements.append(Spacer(1, 6))
            elements.append(Paragraph(f"Item {idx + 1}", styles["Heading4"]))
            elements.append(
                HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#D0D7DE"))
            )
            elements.append(Spacer(1, 6))
            for col in df.columns:
                value = row.get(col)
                text = "" if value is None else str(value)
                elements.append(
                    Paragraph(
                        f"<b><font color='#1F4E79'>{col}:</font></b> {text}",
                        styles["Normal"],
                    )
                )
            elements.append(Spacer(1, 10))

    doc.build(elements)
    buffer.seek(0)
    return FileResponse(
        buffer,
        as_attachment=True,
        filename=filename,
        content_type="application/pdf",
    )


def _build_exercises_catalog():
    qs = (
        Exercise.objects.select_related("exercise_type", "difficulty_level")
        .prefetch_related("equipments")
        .all()
        .order_by("exercise_type__name", "name")
    )
    records = []
    for ex in qs:
        records.append(
            {
                "exercise_type": getattr(ex.exercise_type, "name", str(ex.exercise_type)),
                "name": ex.name,
                "difficulty_level": getattr(ex.difficulty_level, "name", str(ex.difficulty_level)),
                "metabolic_equivalent": ex.metabolic_equivalent,
                "amount_unit": ex.amount_unit,
                "equipments": ", ".join(eq.name for eq in ex.equipments.all()),
                "created_at": _to_naive(ex.created_at),
                "updated_at": _to_naive(ex.updated_at),
            }
        )
    columns = [
        "exercise_type",
        "name",
        "difficulty_level",
        "metabolic_equivalent",
        "amount_unit",
        "equipments",
        "created_at",
        "updated_at",
    ]
    return _df_from_records(records, columns)


def _build_workout_plans_full():
    items_qs = (
        WorkoutItem.objects.select_related(
            "workout_day",
            "workout_day__workout_plan",
            "exercise",
            "exercise__exercise_type",
            "exercise__difficulty_level",
            "workout_day__workout_plan__difficulty_level",
        )
        .prefetch_related("exercise__equipments")
        .order_by(
            "workout_day__workout_plan__id",
            "workout_day__day_number",
            "order",
            "id",
        )
    )
    records = []
    for item in items_qs:
        plan = item.workout_day.workout_plan
        day = item.workout_day
        ex = item.exercise
        records.append(
            {
                "plan_id": plan.id,
                "plan_name": plan.name,
                "plan_difficulty": getattr(
                    plan.difficulty_level, "name", str(plan.difficulty_level)
                ),
                "plan_description": plan.description,
                "day_number": day.day_number,
                "day_description": day.description,
                "item_order": item.order,
                "sets": item.sets,
                "amount": item.amount,
                "amount_unit": ex.amount_unit,
                "exercise_name": ex.name,
                "exercise_type": getattr(ex.exercise_type, "name", str(ex.exercise_type)),
                "exercise_difficulty": getattr(
                    ex.difficulty_level, "name", str(ex.difficulty_level)
                ),
                "metabolic_equivalent": ex.metabolic_equivalent,
                "exercise_equipments": ", ".join(eq.name for eq in ex.equipments.all()),
                "video": ex.video.name if ex.video else "",
                "thumbnail": ex.thumbnail.name if ex.thumbnail else "",
            }
        )
    columns = [
        "plan_id",
        "plan_name",
        "plan_difficulty",
        "plan_description",
        "day_number",
        "day_description",
        "item_order",
        "sets",
        "amount",
        "amount_unit",
        "exercise_name",
        "exercise_type",
        "exercise_difficulty",
        "metabolic_equivalent",
        "exercise_equipments",
        "video",
        "thumbnail",
    ]
    return _df_from_records(records, columns)


def _build_gyms_directory():
    qs = Gym.objects.select_related("address").prefetch_related("equipments").all().order_by("name")
    records = []
    for gym in qs:
        addr = gym.address
        records.append(
            {
                "gym_name": gym.name,
                "street": addr.street,
                "city": addr.city,
                "state": addr.state,
                "postal_code": addr.postal_code,
                "country": addr.country,
                "contact_email": gym.contact_email,
                "contact_phone": gym.contact_phone,
                "equipments": ", ".join(eq.name for eq in gym.equipments.all()),
                "created_at": _to_naive(gym.created_at),
            }
        )
    columns = [
        "gym_name",
        "street",
        "city",
        "state",
        "postal_code",
        "country",
        "contact_email",
        "contact_phone",
        "equipments",
        "created_at",
    ]
    return _df_from_records(records, columns)


def _build_clients_overview():
    qs = (
        ClientProfile.objects.select_related("user", "subscription_plan", "active_workout_plan")
        .all()
        .order_by("user__username")
    )
    records = []
    for client in qs:
        records.append(
            {
                "username": client.user.username,
                "email": client.user.email,
                "is_active": client.user.is_active,
                "verified": client.verified,
                "age": client.age,
                "weight": client.weight,
                "height": client.height,
                "subscription_plan": getattr(client.subscription_plan, "name", None),
                "active_workout_plan": getattr(client.active_workout_plan, "name", None),
                "stripe_customer_id": client.stripe_customer_id,
            }
        )
    columns = [
        "username",
        "email",
        "is_active",
        "verified",
        "age",
        "weight",
        "height",
        "subscription_plan",
        "active_workout_plan",
        "stripe_customer_id",
    ]
    return _df_from_records(records, columns)


def _build_subscriptions_status():
    qs = Subscription.objects.select_related("user").all().order_by("-created_at")
    records = []
    for sub in qs:
        records.append(
            {
                "username": sub.user.username,
                "status": sub.status,
                "price_id": sub.price_id,
                "current_period_start": _to_naive(sub.current_period_start),
                "current_period_end": _to_naive(sub.current_period_end),
                "cancel_at_period_end": sub.cancel_at_period_end,
                "canceled_at": _to_naive(sub.canceled_at),
                "ended_at": _to_naive(sub.ended_at),
                "created_at": _to_naive(sub.created_at),
            }
        )
    columns = [
        "username",
        "status",
        "price_id",
        "current_period_start",
        "current_period_end",
        "cancel_at_period_end",
        "canceled_at",
        "ended_at",
        "created_at",
    ]
    return _df_from_records(records, columns)


def _build_subscription_payments():
    qs = (
        SubscriptionPayment.objects.select_related("subscription", "subscription__user")
        .all()
        .order_by("-paid_at")
    )
    records = []
    for payment in qs:
        records.append(
            {
                "invoice_id": payment.stripe_invoice_id,
                "username": getattr(payment.subscription, "user", None)
                and payment.subscription.user.username,
                "subscription_id": getattr(payment.subscription, "stripe_subscription_id", None),
                "customer_id": payment.customer_id,
                "amount_major": float(Decimal(payment.amount_paid or 0) / Decimal("100")),
                "currency": payment.currency,
                "paid_at": _to_naive(payment.paid_at),
                "created_at": _to_naive(payment.created_at),
            }
        )
    columns = [
        "invoice_id",
        "username",
        "subscription_id",
        "customer_id",
        "amount_major",
        "currency",
        "paid_at",
        "created_at",
    ]
    return _df_from_records(records, columns)


def _build_bug_reports():
    qs = BugReport.objects.select_related("user").all().order_by("-created_at")
    records = []
    for report in qs:
        records.append(
            {
                "title": report.title,
                "user": report.user.username if report.user else "Anonymous",
                "resolved": report.resolved,
                "created_at": _to_naive(report.created_at),
                "resolved_at": _to_naive(report.resolved_at),
            }
        )
    columns = ["title", "user", "resolved", "created_at", "resolved_at"]
    return _df_from_records(records, columns)


def _build_feature_requests():
    qs = NewFeatureRequest.objects.select_related("user").all().order_by("-created_at")
    records = []
    for req in qs:
        records.append(
            {
                "title": req.title,
                "user": req.user.username if req.user else "Anonymous",
                "status": req.status,
                "created_at": _to_naive(req.created_at),
                "implemented_at": _to_naive(req.implemented_at),
            }
        )
    columns = ["title", "user", "status", "created_at", "implemented_at"]
    return _df_from_records(records, columns)


def _build_workout_plan_runs():
    qs = (
        WorkoutPlanRun.objects.select_related("client", "client__user", "workout_plan")
        .annotate(
            total_days=Count("day_logs", distinct=True),
            completed_days=Count("day_logs", filter=Q(day_logs__completed=True), distinct=True),
        )
        .all()
        .order_by("-started_at")
    )
    records = []
    for run in qs:
        total = int(run.total_days or 0)
        completed = int(run.completed_days or 0)
        completion_rate = round((completed / total * 100), 1) if total else 0
        records.append(
            {
                "client": run.client.user.username,
                "workout_plan": run.workout_plan.name,
                "started_at": _to_naive(run.started_at),
                "finished_at": _to_naive(run.finished_at),
                "is_active": run.is_active,
                "total_days": total,
                "completed_days": completed,
                "completion_rate_pct": completion_rate,
            }
        )
    columns = [
        "client",
        "workout_plan",
        "started_at",
        "finished_at",
        "is_active",
        "total_days",
        "completed_days",
        "completion_rate_pct",
    ]
    return _df_from_records(records, columns)


def _build_exercise_usage():
    qs = (
        Exercise.objects.select_related("exercise_type", "difficulty_level")
        .annotate(
            item_count=Count("workoutitem", distinct=True),
            total_sets=Sum("workoutitem__sets"),
            total_amount=Sum("workoutitem__amount"),
            plan_count=Count("workoutitem__workout_day__workout_plan", distinct=True),
        )
        .order_by("-item_count", "name")
    )
    records = []
    for ex in qs:
        records.append(
            {
                "exercise": ex.name,
                "exercise_type": getattr(ex.exercise_type, "name", str(ex.exercise_type)),
                "difficulty": getattr(ex.difficulty_level, "name", str(ex.difficulty_level)),
                "amount_unit": ex.amount_unit,
                "item_count": int(ex.item_count or 0),
                "total_sets": int(ex.total_sets or 0),
                "total_amount": int(ex.total_amount or 0),
                "plan_count": int(ex.plan_count or 0),
            }
        )
    columns = [
        "exercise",
        "exercise_type",
        "difficulty",
        "amount_unit",
        "item_count",
        "total_sets",
        "total_amount",
        "plan_count",
    ]
    return _df_from_records(records, columns)


REPORTS = [
    {
        "key": "exercises_catalog",
        "title": "Exercises Catalog",
        "category": "Programs",
        "description": "Full exercise inventory with types, difficulty, and equipment.",
        "builder": _build_exercises_catalog,
        "sheet": "Exercises",
    },
    {
        "key": "workout_plans_full",
        "title": "Workout Plans (Full)",
        "category": "Programs",
        "description": "Detailed plan breakdown by day and item.",
        "builder": _build_workout_plans_full,
        "sheet": "WorkoutPlans",
    },
    {
        "key": "gyms_directory",
        "title": "Gyms Directory",
        "category": "Operations",
        "description": "All gyms with addresses and contact details.",
        "builder": _build_gyms_directory,
        "sheet": "Gyms",
        "template_enabled": True,
    },
    {
        "key": "clients_overview",
        "title": "Clients Overview",
        "category": "Operations",
        "description": "Client profiles with subscription and plan info.",
        "builder": _build_clients_overview,
        "sheet": "Clients",
        "template_enabled": True,
    },
    {
        "key": "subscriptions_status",
        "title": "Subscriptions Status",
        "category": "Revenue",
        "description": "All subscriptions with lifecycle dates and statuses.",
        "builder": _build_subscriptions_status,
        "sheet": "Subscriptions",
    },
    {
        "key": "subscription_payments",
        "title": "Subscription Payments",
        "category": "Revenue",
        "description": "Payment ledger with invoices and amounts.",
        "builder": _build_subscription_payments,
        "sheet": "Payments",
        "template_enabled": True,
    },
    {
        "key": "bug_reports",
        "title": "Bug Reports",
        "category": "Support",
        "description": "Reported issues with resolution status.",
        "builder": _build_bug_reports,
        "sheet": "BugReports",
    },
    {
        "key": "feature_requests",
        "title": "Feature Requests",
        "category": "Support",
        "description": "Product requests and delivery status.",
        "builder": _build_feature_requests,
        "sheet": "FeatureRequests",
    },
    {
        "key": "workout_plan_runs",
        "title": "Workout Plan Runs",
        "category": "Engagement",
        "description": "Run history with completion metrics.",
        "builder": _build_workout_plan_runs,
        "sheet": "PlanRuns",
    },
    {
        "key": "exercise_usage",
        "title": "Exercise Usage",
        "category": "Engagement",
        "description": "Exercise popularity across plans.",
        "builder": _build_exercise_usage,
        "sheet": "ExerciseUsage",
    },
]


def get_report_catalog():
    categories = {}
    for report in REPORTS:
        categories.setdefault(report["category"], []).append(report)
    return [{"name": category, "reports": items} for category, items in categories.items()]


def _get_report(report_key):
    for report in REPORTS:
        if report["key"] == report_key:
            return report
    return None


def report_xlsx(request, report_key):
    report = _get_report(report_key)
    if not report:
        raise Http404("Report not found")
    df = report["builder"]()
    filename = f"{report_key}.xlsx"
    return _export_xlsx(df, report["sheet"], filename)


def report_pdf(request, report_key):
    report = _get_report(report_key)
    if not report:
        raise Http404("Report not found")
    df = report["builder"]()
    filename = f"{report_key}.pdf"
    return _export_pdf(df, report["title"], filename)


def _build_print_context(report_key):
    report = _get_report(report_key)
    if not report:
        raise Http404("Report not found")
    df = report["builder"]()
    records = df.to_dict(orient="records") if not df.empty else []
    return {
        "title": report["title"],
        "generated_at": timezone.now(),
        "total_records": len(records),
        "records": records,
    }


@login_required
def report_docx(request, report_key, template_id=None):
    report = _get_report(report_key)
    if not report or not report.get("template_enabled"):
        raise Http404("Report template not available")
    ensure_default_templates(report_key)
    template_qs = PrintTemplate.objects.filter(report_key=report_key)
    if template_id:
        template = template_qs.filter(pk=template_id).first()
    else:
        template = template_qs.filter(is_active=True).first()

    if not template:
        raise Http404("No active template for this report")

    context = _build_print_context(report_key)
    doc = DocxTemplate(template.template_file.path)
    doc.render(context)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)

    filename = f"{report_key}-{template.id}.docx"
    return FileResponse(
        output,
        as_attachment=True,
        filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


class PrintTemplateListView(LoginRequiredMixin, ListView):
    model = PrintTemplate
    template_name = "gymApp/print_templates/list.html"
    context_object_name = "templates"

    def get_queryset(self):
        ensure_default_templates()
        return PrintTemplate.objects.all().order_by("report_key", "-updated_at")


class PrintTemplateCreateView(LoginRequiredMixin, CreateView):
    model = PrintTemplate
    form_class = PrintTemplateForm
    template_name = "gymApp/print_templates/form.html"
    success_url = reverse_lazy("gymReports:print_template_list")

    def form_valid(self, form):
        template = form.save(commit=False)
        template.created_by = self.request.user
        template.save()
        self.object = template
        return HttpResponseRedirect(self.get_success_url())


class PrintTemplateUpdateView(LoginRequiredMixin, UpdateView):
    model = PrintTemplate
    form_class = PrintTemplateForm
    template_name = "gymApp/print_templates/form.html"
    success_url = reverse_lazy("gymReports:print_template_list")


class PrintTemplateDeleteView(LoginRequiredMixin, DeleteView):
    model = PrintTemplate
    template_name = "gymApp/print_templates/confirm_delete.html"
    success_url = reverse_lazy("gymReports:print_template_list")


@login_required
def reset_default_template(request, report_key):
    ensure_default_templates(report_key, force=True)
    return HttpResponseRedirect(reverse_lazy("gymReports:print_template_list"))


def reports_index(request):
    ensure_default_templates()
    return render(
        request,
        "gymApp/reports.html",
        {"report_categories": get_report_catalog()},
    )


def exercises_xlsx(request):
    return report_xlsx(request, "exercises_catalog")


def workout_plans_full_report_xlsx(request):
    return report_xlsx(request, "workout_plans_full")
